import os
from datetime import date

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DATA_PATH = "sales_raw.xlsx"
OUTPUT_DIR = "output"


def load_data(path=DATA_PATH):
    """Raw Excel file ko DataFrame me load karo."""
    df = pd.read_excel(path)
    print(f"Loaded rows: {len(df)}")
    return df


def clean_data(df):
    """Step 2: spaces, case, blank, date, duplicates, blank rows."""
    for col in ["Branch", "Payment_Status", "Salesperson"]:
        df[col] = df[col].str.strip()
    df["Branch"] = df["Branch"].str.title()
    df["Payment_Status"] = df["Payment_Status"].str.title()
    df["Salesperson"] = df["Salesperson"].fillna("Unknown")

    df["Date"] = pd.to_datetime(df["Date"], dayfirst=True,
                                errors="coerce", format="mixed")
    df = df.drop_duplicates()
    df = df.dropna(subset=["Date", "Amount"]).reset_index(drop=True)
    print(f"Clean rows: {len(df)}")
    return df


def flag_outliers(df):
    """Step 3: unit price median se 3 guna zyada ho to outlier."""
    unit = df["Amount"] / df["Qty"]
    median = unit.groupby(df["Product"]).transform("median")
    df["Is_Outlier"] = unit > 3 * median
    print(f"Outliers flagged: {int(df['Is_Outlier'].sum())}")
    return df


def make_tables(df):
    """Step 4: KPIs aur summary tables."""
    df["Month"] = df["Date"].dt.to_period("M").astype(str)

    kpi = pd.DataFrame({
        "KPI": ["Total Sales", "Total Orders", "Avg Order Value"],
        "Value": [df["Amount"].sum(), len(df), round(df["Amount"].mean(), 2)],
    })
    by_branch = (df.groupby("Branch")["Amount"].sum()
                   .sort_values(ascending=False).reset_index())
    by_month = df.groupby("Month")["Amount"].sum().reset_index()
    top_products = (df.groupby("Product")["Amount"].sum()
                      .nlargest(5).reset_index())
    pending = df[df["Payment_Status"] == "Pending"][
        ["Date", "Branch", "Product", "Salesperson", "Qty", "Amount"]].copy()
    pending["Date"] = pending["Date"].dt.date

    return {"Summary": kpi, "Branch": by_branch, "Monthly": by_month,
            "TopProducts": top_products, "Pending": pending}


def add_chart(ws, chart, title, anchor="E2"):
    chart.title = title
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    chart.width, chart.height = 18, 9
    ws.add_chart(chart, anchor)


def build_report(tables):
    """Step 5: Excel report with formatting aur charts."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    out = f"{OUTPUT_DIR}/MIS_Report_{date.today()}.xlsx"

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name, index=False)

    wb = load_workbook(out)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for col in ws.columns:
            longest = max(len(str(c.value)) if c.value is not None else 0
                          for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = longest + 4
        for col in ws.iter_cols(min_row=2):
            if ws.cell(row=1, column=col[0].column).value in ("Amount", "Value"):
                for c in col:
                    c.number_format = "#,##0"

    wb["Summary"]["B4"].number_format = "#,##0.00"   # Avg Order Value
    add_chart(wb["Branch"], BarChart(), "Branch-wise Sales")
    add_chart(wb["Monthly"], LineChart(), "Month-wise Sales")

    wb.save(out)
    return out


def main():
    df = load_data()
    df = clean_data(df)
    df = flag_outliers(df)
    tables = make_tables(df)
    out = build_report(tables)
    print("Report ready:", out)


if __name__ == "__main__":
    main()

    